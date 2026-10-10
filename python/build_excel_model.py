"""Build excel/healthcare_opportunity_model.xlsx (live formulas, visible assumptions).

    python python/build_excel_model.py      # run after the notebook has exported data/processed/*.csv

Sheets
  Summary          headline KPIs + scenario results (formulas) + chart
  Assumptions      every input the model uses (blue = editable)
  Hospital_Screen  1,996 hospitals, largest severity shortfall first (data from SQL view)
  Scenarios        scope x capture-rate grid, sensitivity table, and Python cross-checks
  v1_Reconciliation  what was wrong in the first version of this model, and the corrected numbers
"""
import sqlite3
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "excel" / "healthcare_opportunity_model.xlsx"

INK, ACCENT, GREY = "1F2937", "1D4E89", "6B7280"
H_FILL = PatternFill("solid", fgColor=ACCENT)
IN_FILL = PatternFill("solid", fgColor="FFF8DC")
H_FONT = Font(bold=True, color="FFFFFF")
IN_FONT = Font(color="0000FF", bold=True)              # blue = input
THIN = Side(style="thin", color="D1D5DB")
USD0, USD_M, PCT1, INT = '"$"#,##0', '"$"#,##0.0,,"M"', "0.0%", "#,##0"


def header(ws, row, labels, col=1):
    for i, lab in enumerate(labels):
        c = ws.cell(row=row, column=col + i, value=lab)
        c.font, c.fill = H_FONT, H_FILL
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def widths(ws, w):
    for i, x in enumerate(w, 1):
        ws.column_dimensions[get_column_letter(i)].width = x


def main() -> None:
    scr = pd.read_csv(ROOT / "data" / "processed" / "excel_severity_inputs.csv", dtype={"provider_id": str})
    con = sqlite3.connect(ROOT / "data" / "healthcare.db")
    k = pd.read_sql("SELECT SUM(total_discharges) d, SUM(total_charges) ch, SUM(total_payment) pay, "
                    "SUM(total_medicare_payment) mc, SUM(charge_gap) gap FROM v_fact", con).iloc[0]
    n = len(scr)
    last = n + 1                                                   # last data row in Hospital_Screen

    wb = Workbook()
    # ------------------------------------------------------------------ Assumptions
    wa = wb.active; wa.title = "Assumptions"
    wa["A1"] = "Assumptions (blue cells are inputs - change them and every sheet recalculates)"; wa["A1"].font = Font(bold=True, size=13, color=INK)
    rows = [
        ("Review threshold: minimum shortfall vs peers (% of expected payment)", 0.02, PCT1,
         "Only hospitals at least this far below peers' severity mix are counted. Filters statistical noise."),
        ("Capture rate - Low", 0.10, PCT1, "Share of the shortfall realistically supported by clinical documentation review"),
        ("Capture rate - Base", 0.25, PCT1, ""),
        ("Capture rate - High", 0.50, PCT1, ""),
        ("Scope A: top N hospitals reviewed", 25, INT, "Pilot wave: largest shortfalls first"),
        ("Scope B: top N hospitals reviewed", 100, INT, "Programme wave"),
        ("Scope C: top N hospitals reviewed", 10000, INT, "All eligible hospitals (N larger than the list)"),
    ]
    header(wa, 3, ["Input", "Value", "Notes"])
    for i, (lab, val, fmt, note) in enumerate(rows, start=4):
        wa.cell(row=i, column=1, value=lab)
        c = wa.cell(row=i, column=2, value=val); c.number_format = fmt; c.font = IN_FONT; c.fill = IN_FILL
        wa.cell(row=i, column=3, value=note).font = Font(color=GREY)
    wa["A12"] = "Fixed facts (from the SQL views, not editable)"; wa["A12"].font = Font(bold=True, color=INK)
    facts = [("Total charges, all hospitals", k.ch, USD0), ("Total payments, all payers", k.pay, USD0),
             ("   of which Medicare", k.mc, USD0), ("Charge-to-payment gap (descriptive, NOT recoverable)", k.gap, USD0),
             ("Discharges", k.d, INT)]
    for i, (lab, val, fmt) in enumerate(facts, start=13):
        wa.cell(row=i, column=1, value=lab); wa.cell(row=i, column=2, value=float(val)).number_format = fmt
    wa["A19"] = ("Source: CMS Medicare Inpatient Hospitals by Provider and Service, data year 2023 (fee-for-service Medicare only; "
                 "cells of 10 or fewer discharges suppressed).")
    wa["A19"].font = Font(italic=True, color=GREY)
    widths(wa, [66, 18, 90])
    THR, CL, CB, CH, NA, NB, NC = (f"Assumptions!$B${r}" for r in range(4, 11))

    # ------------------------------------------------------------------ Hospital_Screen
    wh = wb.create_sheet("Hospital_Screen")
    cols = ["Provider ID", "Hospital", "State", "Setting", "Tiered discharges", "Top-tier share", "Peer expected share",
            "Expected payment ($)", "Severity shortfall ($)", "Shortfall %", "Eligible (1/0)", "Eligible shortfall ($)",
            "Eligible rank"]
    header(wh, 1, cols)
    for r, row in enumerate(scr.itertuples(index=False), start=2):
        vals = [row.provider_id, row.provider_name, row.state_code, row.setting, row.tiered_discharges, row.top_tier_share,
                row.peer_expected_top_tier_share, row.expected_value, row.severity_shortfall, row.shortfall_pct]
        for c, v in enumerate(vals, 1):
            wh.cell(row=r, column=c, value=v)
        wh.cell(row=r, column=11, value=f"=IF(AND(I{r}>0,J{r}>={THR}),1,0)")
        wh.cell(row=r, column=12, value=f"=I{r}*K{r}")
        wh.cell(row=r, column=13, value=f"=SUM($K$2:K{r})")           # list is sorted by shortfall desc
        for c, fmt in [(5, INT), (6, PCT1), (7, PCT1), (8, USD0), (9, USD0), (10, PCT1), (12, USD0)]:
            wh.cell(row=r, column=c).number_format = fmt
    wh.freeze_panes = "C2"; wh.auto_filter.ref = f"A1:M{last}"
    widths(wh, [11, 44, 7, 15, 12, 11, 12, 16, 16, 10, 10, 16, 10])
    wh.row_dimensions[1].height = 32
    SH, SJ, SL, SM = (f"Hospital_Screen!${c}$2:${c}${last}" for c in "IJLM")

    # ------------------------------------------------------------------ Scenarios
    ws = wb.create_sheet("Scenarios")
    ws["A1"] = "Documentation-capture opportunity: annual payment effect, all payers on Medicare cases"; ws["A1"].font = Font(bold=True, size=13, color=INK)
    ws["A2"] = "Eligible shortfall x capture rate. A screen for clinical review - not a revenue forecast."; ws["A2"].font = Font(italic=True, color=GREY)
    header(ws, 4, ["Scope", "Hospitals in scope", "Eligible shortfall ($)", "Low", "Base", "High"])
    for r, (lab, nref) in enumerate([("A. Pilot", NA), ("B. Programme", NB), ("C. All eligible", NC)], start=5):
        ws.cell(row=r, column=1, value=f'="{lab} (top "&TEXT(MIN({nref},MAX({SM})),"#,##0")&" hospitals)"')
        ws.cell(row=r, column=2, value=f"=MIN({nref},MAX({SM}))").number_format = INT
        ws.cell(row=r, column=3, value=f'=SUMIFS({SL},{SM},"<="&{nref})').number_format = USD0
        for c, cap in zip((4, 5, 6), (CL, CB, CH)):
            ws.cell(row=r, column=c, value=f"=$C{r}*{cap}").number_format = USD0
    ws["A9"] = "Eligible hospitals / total screened"; ws["B9"] = f"=MAX({SM})"; ws["C9"] = f"=COUNT({SH})"
    ws["B9"].number_format = ws["C9"].number_format = INT
    ws["A10"] = "Eligible shortfall as % of expected payments on tiered DRGs"
    ws["C10"] = f"=SUMIFS({SL},{SM},\">0\")/SUM(Hospital_Screen!$H$2:$H${last})"; ws["C10"].number_format = '0.00%'

    ws["A13"] = "Sensitivity: all hospitals at or above the threshold, by capture rate"; ws["A13"].font = Font(bold=True, color=INK)
    thr_list = [0.0, 0.01, 0.02, 0.03, 0.05]
    header(ws, 14, ["Capture rate \\ review threshold"] + [None] * len(thr_list))
    for j, t in enumerate(thr_list, start=2):
        c = ws.cell(row=14, column=j, value=t); c.number_format = '0%'; c.font = H_FONT; c.fill = H_FILL
    for i, cap in enumerate([0.05, 0.10, 0.25, 0.50], start=15):
        c = ws.cell(row=i, column=1, value=cap); c.number_format = '0%'; c.font = Font(bold=True)
        for j in range(2, 2 + len(thr_list)):
            col = get_column_letter(j)
            ws.cell(row=i, column=j, value=f'=SUMIFS({SH},{SJ},">="&{col}$14,{SH},">0")*$A{i}').number_format = USD_M

    ws["A21"] = "Cross-check: same figures computed in Python from the SQL view (should equal the Base column and the 2% / 25% sensitivity cell)"
    ws["A21"].font = Font(bold=True, color=INK)
    elig = scr[(scr.severity_shortfall > 0) & (scr.shortfall_pct >= 0.02)].sort_values("severity_shortfall", ascending=False)
    py = {"A": elig.severity_shortfall.head(25).sum() * .25, "B": elig.severity_shortfall.head(100).sum() * .25,
          "C": elig.severity_shortfall.sum() * .25}
    header(ws, 22, ["Scope", "Python (Base 25%)", "Excel (Base)", "Difference"])
    for r, (key, xr) in enumerate([("A", 5), ("B", 6), ("C", 7)], start=23):
        ws.cell(row=r, column=1, value=key)
        ws.cell(row=r, column=2, value=float(py[key])).number_format = USD0
        ws.cell(row=r, column=3, value=f"=E{xr}").number_format = USD0
        ws.cell(row=r, column=4, value=f"=C{r}-B{r}").number_format = USD0
    widths(ws, [52, 20, 22, 18, 18, 18]); ws.freeze_panes = "A5"

    # ------------------------------------------------------------------ v1 reconciliation
    wr = wb.create_sheet("v1_Reconciliation")
    wr["A1"] = "What changed from v1 of this model, and why"; wr["A1"].font = Font(bold=True, size=13, color=INK)
    header(wr, 3, ["v1 figure", "v1 value", "Corrected / retired", "Cause"])
    rec = [
        ("Total 'revenue gap'", 717.55e9, k.gap,
         "v1 summed an entire pivot-table column, including its Grand Total row, so every dollar counted twice "
         "($358.8B on Medicare payment). v2 uses all-payer payment: $343.8B."),
        ("'Top DRG optimisation' (8%)", 57.4e9, 8.99e9,
         "v1 applied 8% to the whole gap, not the top 10 DRGs (31% of the gap). Correct v1 logic = 8% x $112.5B."),
        ("'Top 20 providers' (10%)", 0.236e9, 4.65e9,
         "v1 summed the first 20 rows of an unsorted pivot, not the largest 20. Correct v1 logic = 10% x $46.5B."),
        ("'5% efficiency gain'", 35.9e9, 17.9e9, "Halved once the doubled total is fixed (5% x $358.8B)."),
        ("'1% efficiency gain'", 7.18e9, 3.59e9, "Same doubling."),
    ]
    for r, (a, b, c, d) in enumerate(rec, start=4):
        wr.cell(row=r, column=1, value=a)
        wr.cell(row=r, column=2, value=b).number_format = USD0
        wr.cell(row=r, column=3, value=float(c)).number_format = USD0
        wr.cell(row=r, column=4, value=d).alignment = Alignment(wrap_text=True, vertical="top")
        wr.row_dimensions[r].height = 48
    wr["A10"] = ("More important than the arithmetic: charges are list prices, not expected revenue, and Medicare pays a fixed DRG rate "
                 "regardless of charges. v1's 'recovery' scenarios therefore described money that was never collectable. v2 replaces "
                 "them with the one lever hospitals control - documentation of case severity.")
    wr["A10"].alignment = Alignment(wrap_text=True, vertical="top"); wr.merge_cells("A10:D10"); wr.row_dimensions[10].height = 62
    widths(wr, [30, 18, 20, 100])

    # ------------------------------------------------------------------ Summary (first sheet)
    wsum = wb.create_sheet("Summary", 0)
    wsum["A1"] = "Medicare Inpatient: Pricing, Payment & Documentation Opportunity (DY2023)"; wsum["A1"].font = Font(bold=True, size=14, color=INK)
    header(wsum, 3, ["Headline KPI", "Value"])
    kp = [("Total charges (list price)", "=Assumptions!B13", USD0), ("Total payments (all payers)", "=Assumptions!B14", USD0),
          ("Payment-to-charge ratio", "=Assumptions!B14/Assumptions!B13", PCT1),
          ("Charge-to-payment gap (descriptive only)", "=Assumptions!B16", USD0),
          ("Hospitals screened for severity mix", "=Scenarios!C9", INT),
          ("Eligible for documentation review", "=Scenarios!B9", INT)]
    for r, (lab, f, fmt) in enumerate(kp, start=4):
        wsum.cell(row=r, column=1, value=lab); wsum.cell(row=r, column=2, value=f).number_format = fmt
    header(wsum, 11, ["Documentation-capture opportunity", "Low", "Base", "High"])
    for r, xr in enumerate((5, 6, 7), start=12):
        wsum.cell(row=r, column=1, value=f"=Scenarios!A{xr}")
        for c, col in zip((2, 3, 4), "DEF"):
            wsum.cell(row=r, column=c, value=f"=Scenarios!{col}{xr}").number_format = USD_M
    wsum["A16"] = "Read the Assumptions sheet before quoting any number. See v1_Reconciliation for corrections to the earlier model."
    wsum["A16"].font = Font(italic=True, color=GREY)
    ch = BarChart(); ch.type = "col"; ch.title = "Annual payment effect by scope (Base capture)"; ch.style = 10
    ch.add_data(Reference(wsum, min_col=3, min_row=11, max_row=14), titles_from_data=True)
    ch.set_categories(Reference(wsum, min_col=1, min_row=12, max_row=14))
    ch.series[0].graphicalProperties.solidFill = ACCENT; ch.legend = None
    ch.y_axis.numFmt = '"$"#,##0,,"M"'; ch.height, ch.width = 7.5, 15
    wsum.add_chart(ch, "F3")
    widths(wsum, [52, 20, 18, 18])

    OUT.parent.mkdir(exist_ok=True)
    wb.save(OUT)
    print("saved", OUT, f"({n:,} hospitals)")
    print({k_: round(v / 1e6, 1) for k_, v in py.items()}, "Python base-case $M by scope")


if __name__ == "__main__":
    main()
